"""
Excel 导出工具
将 JSON 数据导出为格式化的 Excel 文件
"""
import json
import pandas as pd
from typing import Dict, Any, List
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class ExcelExporter:
    """Excel 导出工具类"""
    
    @staticmethod
    def export_to_excel(data: str, filename: str, columns: List[str] = None) -> Dict[str, Any]:
        """
        将 JSON 数据导出为 Excel 文件
        
        Args:
            data: JSON 格式的字符串数据
            filename: 输出的 Excel 文件名
            columns: 列名列表，如果为 None 则自动推断
            
        Returns:
            dict: 包含成功状态和文件路径的字典
            
        Raises:
            Exception: 导出过程中的错误
        """
        try:
            # 解析 JSON 数据
            if isinstance(data, str):
                test_cases = json.loads(data)
            else:
                test_cases = data
            
            if not test_cases:
                raise ValueError("测试用例数据为空")
            
            # 处理 steps 字段（列表转字符串）
            for case in test_cases:
                if 'steps' in case and isinstance(case['steps'], list):
                    case['steps'] = '\n'.join([f"{i+1}. {step}" for i, step in enumerate(case['steps'])])
            
            # 创建 DataFrame
            df = pd.DataFrame(test_cases)
            
            # 如果指定了列名，重新排序
            if columns:
                # 只保留存在的列
                columns = [col for col in columns if col in df.columns]
                df = df[columns]
            
            # 默认列名映射（中文）
            column_mapping = {
                'title': '测试用例标题',
                'description': '描述',
                'precondition': '前置条件',
                'steps': '测试步骤',
                'expected_result': '预期结果',
                'actual_result': '实际结果',
                'pass_fail': '通过/失败'
            }
            
            # 重命名列
            df = df.rename(columns=column_mapping)
            
            # 导出到 Excel
            df.to_excel(filename, index=False, engine='openpyxl')
            
            # 美化 Excel 格式
            ExcelExporter._beautify_excel(filename)
            
            return {
                "success": True,
                "filename": filename,
                "rows": len(test_cases),
                "message": f"成功导出 {len(test_cases)} 条测试用例到 {filename}"
            }
            
        except json.JSONDecodeError as e:
            raise Exception(f"JSON 解析失败: {str(e)}")
        except Exception as e:
            raise Exception(f"导出 Excel 失败: {str(e)}")
    
    @staticmethod
    def _beautify_excel(filename: str):
        """
        美化 Excel 格式
        - 设置列宽
        - 添加边框
        - 设置标题样式
        - 文本对齐
        
        Args:
            filename: Excel 文件路径
        """
        try:
            wb = load_workbook(filename)
            ws = wb.active
            
            # 标题行样式
            header_font = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # 边框样式
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # 内容对齐
            content_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            
            # 设置标题行样式
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # 设置内容样式和列宽
            column_widths = {
                'A': 30,  # 测试用例标题
                'B': 35,  # 描述
                'C': 25,  # 前置条件
                'D': 50,  # 测试步骤
                'E': 30,  # 预期结果
                'F': 20,  # 实际结果
                'G': 15   # 通过/失败
            }
            
            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width
            
            # 设置内容单元格样式
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = content_alignment
                    cell.border = thin_border
            
            # 设置行高
            for row in range(2, ws.max_row + 1):
                ws.row_dimensions[row].height = None  # 自动调整
            
            wb.save(filename)
            
        except Exception as e:
            # 美化失败不影响导出，只记录错误
            print(f"Warning: Excel 美化失败: {str(e)}")
